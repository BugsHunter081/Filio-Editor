import io
import math
import os
import re
import zipfile
from pathlib import Path
from typing import Any

import cv2
import numpy as np
import pymupdf
from PIL import Image, ImageOps
from pypdf import PdfReader, PdfWriter
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas as pdf_canvas


OUTPUT_DIR = Path(__file__).resolve().parent.parent.parent / "storage" / "outputs"
PREVIEW_DIR = Path(__file__).resolve().parent.parent.parent / "storage" / "previews"
OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
PREVIEW_DIR.mkdir(parents=True, exist_ok=True)


def _open(path: str) -> pymupdf.Document:
    return pymupdf.open(path)


def _save(doc: pymupdf.Document, dest: Path) -> Path:
    dest.parent.mkdir(parents=True, exist_ok=True)
    doc.save(dest, garbage=4, deflate=True, clean=True)
    doc.close()
    return dest


def page_count(path: str) -> int:
    doc = _open(path)
    n = doc.page_count
    doc.close()
    return n


def render_preview(path: str, dest: Path, max_pages: int = 12, dpi: int = 72) -> list[str]:
    doc = _open(path)
    dest.mkdir(parents=True, exist_ok=True)
    names = []
    for i in range(min(doc.page_count, max_pages)):
        pix = doc[i].get_pixmap(dpi=dpi, alpha=False)
        name = f"page-{i + 1}.jpg"
        pix.save(str(dest / name))
        names.append(name)
    doc.close()
    return names


def parse_ranges(spec: str, total: int) -> list[int]:
    if not spec or not spec.strip():
        return list(range(total))
    pages: set[int] = set()
    for part in spec.split(","):
        part = part.strip()
        if not part:
            continue
        if "-" in part:
            a, b = part.split("-", 1)
            start = max(1, int(a))
            end = min(total, int(b))
            pages.update(range(start - 1, end))
        else:
            n = int(part)
            if 1 <= n <= total:
                pages.add(n - 1)
    return sorted(pages)


def merge(paths: list[str], dest: Path, options: dict) -> Path:
    out = pymupdf.open()
    for p in paths:
        src = pymupdf.open(p)
        out.insert_pdf(src)
        src.close()
    if options.get("cover_title"):
        page = out.new_page(pno=0, width=595, height=842)
        rect = pymupdf.Rect(80, 320, 515, 420)
        page.insert_textbox(rect, options["cover_title"], fontsize=28, fontname="helv", align=1)
    return _save(out, dest)


def split_by_pages(path: str, dest_dir: Path, options: dict) -> list[Path]:
    doc = _open(path)
    ranges = options.get("ranges") or "1"
    parts = []
    dest_dir.mkdir(parents=True, exist_ok=True)
    chunks = [c.strip() for c in str(ranges).split(";") if c.strip()] or ["1"]
    for i, chunk in enumerate(chunks, 1):
        pages = parse_ranges(chunk, doc.page_count)
        if not pages:
            continue
        out = pymupdf.open()
        for pg in pages:
            out.insert_pdf(doc, from_page=pg, to_page=pg)
        p = dest_dir / f"split-{i}.pdf"
        out.save(p, garbage=4, deflate=True)
        out.close()
        parts.append(p)
    doc.close()
    return parts


def split_in_half(path: str, dest_dir: Path) -> list[Path]:
    doc = _open(path)
    mid = max(1, doc.page_count // 2)
    dest_dir.mkdir(parents=True, exist_ok=True)
    a = pymupdf.open()
    b = pymupdf.open()
    a.insert_pdf(doc, from_page=0, to_page=mid - 1)
    if mid < doc.page_count:
        b.insert_pdf(doc, from_page=mid, to_page=doc.page_count - 1)
    pa, pb = dest_dir / "half-1.pdf", dest_dir / "half-2.pdf"
    a.save(pa)
    b.save(pb)
    a.close()
    b.close()
    doc.close()
    return [pa, pb]


def split_by_size(path: str, dest_dir: Path, max_mb: float = 2.0) -> list[Path]:
    doc = _open(path)
    dest_dir.mkdir(parents=True, exist_ok=True)
    limit = max_mb * 1024 * 1024
    parts = []
    current = pymupdf.open()
    idx = 1
    for i in range(doc.page_count):
        current.insert_pdf(doc, from_page=i, to_page=i)
        buf = current.tobytes(deflate=True)
        if len(buf) >= limit and current.page_count > 1:
            current.delete_page(current.page_count - 1)
            p = dest_dir / f"size-{idx}.pdf"
            current.save(p)
            current.close()
            parts.append(p)
            idx += 1
            current = pymupdf.open()
            current.insert_pdf(doc, from_page=i, to_page=i)
    if current.page_count:
        p = dest_dir / f"size-{idx}.pdf"
        current.save(p)
        parts.append(p)
    current.close()
    doc.close()
    return parts


def extract_pages(path: str, dest: Path, options: dict) -> Path:
    doc = _open(path)
    pages = parse_ranges(options.get("ranges", "1"), doc.page_count)
    out = pymupdf.open()
    for pg in pages:
        out.insert_pdf(doc, from_page=pg, to_page=pg)
    doc.close()
    return _save(out, dest)


def delete_pages(path: str, dest: Path, options: dict) -> Path:
    doc = _open(path)
    pages = set(parse_ranges(options.get("ranges", ""), doc.page_count))
    keep = [i for i in range(doc.page_count) if i not in pages]
    out = pymupdf.open()
    for i in keep:
        out.insert_pdf(doc, from_page=i, to_page=i)
    doc.close()
    return _save(out, dest)


def rotate_pages(path: str, dest: Path, options: dict) -> Path:
    doc = _open(path)
    angle = int(options.get("angle", 90))
    pages = parse_ranges(options.get("ranges", ""), doc.page_count)
    for i in pages:
        doc[i].set_rotation((doc[i].rotation + angle) % 360)
    return _save(doc, dest)


def crop_pages(path: str, dest: Path, options: dict) -> Path:
    doc = _open(path)
    margin = float(options.get("margin", 36))
    pages = parse_ranges(options.get("ranges", ""), doc.page_count)
    for i in pages:
        page = doc[i]
        r = page.rect
        page.set_cropbox(pymupdf.Rect(r.x0 + margin, r.y0 + margin, r.x1 - margin, r.y1 - margin))
    return _save(doc, dest)


def resize_pages(path: str, dest: Path, options: dict) -> Path:
    sizes = {
        "a4": (595, 842),
        "letter": (612, 792),
        "legal": (612, 1008),
        "a5": (420, 595),
    }
    target = sizes.get(str(options.get("page_size", "a4")).lower(), (595, 842))
    src = _open(path)
    out = pymupdf.open()
    pages = parse_ranges(options.get("ranges", ""), src.page_count)
    selected = set(pages)
    for i in range(src.page_count):
        if i not in selected:
            out.insert_pdf(src, from_page=i, to_page=i)
            continue
        page = out.new_page(width=target[0], height=target[1])
        page.show_pdf_page(page.rect, src, i)
    src.close()
    return _save(out, dest)


def organize(path: str, dest: Path, options: dict) -> Path:
    order = options.get("order") or []
    doc = _open(path)
    out = pymupdf.open()
    if not order:
        order = list(range(doc.page_count))
    for i in order:
        i = int(i)
        if 0 <= i < doc.page_count:
            out.insert_pdf(doc, from_page=i, to_page=i)
    doc.close()
    return _save(out, dest)


def compress(path: str, dest: Path, options: dict) -> Path:
    level = str(options.get("level", "medium"))
    dpi_map = {"low": 150, "medium": 110, "high": 72}
    quality_map = {"low": 80, "medium": 60, "high": 40}
    dpi = dpi_map.get(level, 110)
    quality = quality_map.get(level, 60)
    src = _open(path)
    out = pymupdf.open()
    for i in range(src.page_count):
        page = src[i]
        pix = page.get_pixmap(dpi=dpi, alpha=False)
        img = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
        buf = io.BytesIO()
        img.save(buf, format="JPEG", quality=quality, optimize=True)
        img_pdf = pymupdf.open("pdf", pymupdf.open("jpeg", buf.getvalue()).convert_to_pdf())
        new_page = out.new_page(width=page.rect.width, height=page.rect.height)
        new_page.show_pdf_page(new_page.rect, img_pdf, 0)
        img_pdf.close()
    src.close()
    return _save(out, dest)


def ocr_pdf(path: str, dest: Path, options: dict) -> Path:
    import pytesseract

    lang = options.get("language", "eng")
    src = _open(path)
    out = pymupdf.open()
    pages = parse_ranges(options.get("ranges", ""), src.page_count)
    selected = set(pages)
    for i in range(src.page_count):
        page = src[i]
        if i not in selected:
            out.insert_pdf(src, from_page=i, to_page=i)
            continue
        pix = page.get_pixmap(dpi=200, alpha=False)
        img = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
        text = pytesseract.image_to_string(img, lang=lang)
        new_page = out.new_page(width=page.rect.width, height=page.rect.height)
        new_page.show_pdf_page(new_page.rect, src, i)
        if text.strip():
            new_page.insert_textbox(
                new_page.rect,
                text,
                fontsize=0.1,
                render_mode=3,
            )
    src.close()
    return _save(out, dest)


def deskew(path: str, dest: Path, options: dict) -> Path:
    src = _open(path)
    out = pymupdf.open()
    pages = parse_ranges(options.get("ranges", ""), src.page_count)
    selected = set(pages)
    for i in range(src.page_count):
        page = src[i]
        if i not in selected:
            out.insert_pdf(src, from_page=i, to_page=i)
            continue
        pix = page.get_pixmap(dpi=150, alpha=False)
        arr = np.frombuffer(pix.samples, dtype=np.uint8).reshape(pix.height, pix.width, pix.n)
        gray = cv2.cvtColor(arr, cv2.COLOR_RGB2GRAY)
        gray = cv2.bitwise_not(gray)
        coords = np.column_stack(np.where(gray > 0))
        angle = 0.0
        if len(coords) > 20:
            rect = cv2.minAreaRect(coords)
            angle = rect[-1]
            if angle < -45:
                angle = 90 + angle
            if angle > 45:
                angle = angle - 90
        mat = cv2.getRotationMatrix2D((pix.width / 2, pix.height / 2), angle, 1.0)
        rotated = cv2.warpAffine(arr, mat, (pix.width, pix.height), flags=cv2.INTER_CUBIC, borderValue=(255, 255, 255))
        img = Image.fromarray(rotated)
        buf = io.BytesIO()
        img.save(buf, format="JPEG", quality=90)
        img_doc = pymupdf.open("jpeg", buf.getvalue())
        pdfbytes = img_doc.convert_to_pdf()
        img_doc.close()
        img_pdf = pymupdf.open("pdf", pdfbytes)
        new_page = out.new_page(width=page.rect.width, height=page.rect.height)
        new_page.show_pdf_page(new_page.rect, img_pdf, 0)
        img_pdf.close()
    src.close()
    return _save(out, dest)


def protect(path: str, dest: Path, options: dict) -> Path:
    password = options.get("password") or "folio"
    writer = PdfWriter()
    reader = PdfReader(path)
    for page in reader.pages:
        writer.add_page(page)
    perms = 0
    if options.get("allow_print"):
        perms |= 4
    if options.get("allow_copy"):
        perms |= 16
    if options.get("allow_edit"):
        perms |= 8
    writer.encrypt(user_password=password, owner_password=password + "-owner", permissions_flag=perms)
    dest.parent.mkdir(parents=True, exist_ok=True)
    with open(dest, "wb") as f:
        writer.write(f)
    return dest


def unlock(path: str, dest: Path, options: dict) -> Path:
    password = options.get("password") or ""
    try:
        reader = PdfReader(path)
        if reader.is_encrypted:
            ok = reader.decrypt(password)
            if not ok:
                raise ValueError("Incorrect password")
        writer = PdfWriter()
        for page in reader.pages:
            writer.add_page(page)
        dest.parent.mkdir(parents=True, exist_ok=True)
        with open(dest, "wb") as f:
            writer.write(f)
        return dest
    except Exception:
        doc = pymupdf.open(path, password=password)
        if doc.is_encrypted:
            if not doc.authenticate(password):
                raise ValueError("Incorrect password")
        return _save(doc, dest)


def watermark(path: str, dest: Path, options: dict) -> Path:
    text = options.get("text") or "CONFIDENTIAL"
    opacity = float(options.get("opacity", 0.18))
    rotation = float(options.get("rotation", 45))
    size = float(options.get("size", 48))
    doc = _open(path)
    pages = parse_ranges(options.get("ranges", ""), doc.page_count)
    for i in pages:
        page = doc[i]
        r = page.rect
        shape = page.new_shape()
        shape.insert_text(
            pymupdf.Point(r.width * 0.18, r.height * 0.55),
            text,
            fontsize=size,
            rotate=int(rotation),
            color=(0.4, 0.4, 0.45),
            fill_opacity=opacity,
            overlay=True,
        )
        shape.commit()
    return _save(doc, dest)


def flatten_pdf(path: str, dest: Path, options: dict) -> Path:
    doc = _open(path)
    try:
        doc.bake()
    except Exception:
        for page in doc:
            pix = page.get_pixmap(dpi=120, alpha=False)
            img_doc = pymupdf.open("png", pix.tobytes("png"))
            pdfbytes = img_doc.convert_to_pdf()
            img_doc.close()
    return _save(doc, dest)


def edit_metadata(path: str, dest: Path, options: dict) -> Path:
    doc = _open(path)
    meta = doc.metadata or {}
    for key in ("title", "author", "subject", "keywords", "creator", "producer"):
        if options.get(key) is not None:
            meta[key] = options[key]
    doc.set_metadata(meta)
    return _save(doc, dest)


def create_bookmarks(path: str, dest: Path, options: dict) -> Path:
    doc = _open(path)
    toc = []
    items = options.get("bookmarks") or []
    for item in items:
        title = item.get("title") or f"Page {item.get('page', 1)}"
        page = max(1, int(item.get("page", 1)))
        level = max(1, int(item.get("level", 1)))
        toc.append([level, title, page])
    if toc:
        doc.set_toc(toc)
    return _save(doc, dest)


def header_footer(path: str, dest: Path, options: dict) -> Path:
    header = options.get("header", "")
    footer = options.get("footer", "")
    numbering = bool(options.get("page_numbers", True))
    doc = _open(path)
    pages = parse_ranges(options.get("ranges", ""), doc.page_count)
    for i in pages:
        page = doc[i]
        r = page.rect
        if header:
            page.insert_text((r.x0 + 36, r.y0 + 24), header, fontsize=10, color=(0.2, 0.2, 0.25))
        if footer or numbering:
            text = footer
            if numbering:
                text = (footer + "  " if footer else "") + f"{i + 1} / {doc.page_count}"
            page.insert_text((r.x0 + 36, r.y1 - 24), text.strip(), fontsize=10, color=(0.3, 0.3, 0.35))
    return _save(doc, dest)


def remove_annotations(path: str, dest: Path, options: dict) -> Path:
    doc = _open(path)
    pages = parse_ranges(options.get("ranges", ""), doc.page_count)
    for i in pages:
        page = doc[i]
        for annot in list(page.annots() or []):
            page.delete_annot(annot)
    return _save(doc, dest)


def grayscale(path: str, dest: Path, options: dict) -> Path:
    src = _open(path)
    out = pymupdf.open()
    for i in range(src.page_count):
        page = src[i]
        pix = page.get_pixmap(dpi=130, colorspace=pymupdf.csGRAY, alpha=False)
        img_doc = pymupdf.open("png", pix.tobytes("png"))
        pdfbytes = img_doc.convert_to_pdf()
        img_doc.close()
        img_pdf = pymupdf.open("pdf", pdfbytes)
        npage = out.new_page(width=page.rect.width, height=page.rect.height)
        npage.show_pdf_page(npage.rect, img_pdf, 0)
        img_pdf.close()
    src.close()
    return _save(out, dest)


def flip_pages(path: str, dest: Path, options: dict) -> Path:
    src = _open(path)
    out = pymupdf.open()
    axis = options.get("axis", "horizontal")
    for i in range(src.page_count):
        page = src[i]
        pix = page.get_pixmap(dpi=130, alpha=False)
        img = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
        img = ImageOps.mirror(img) if axis == "horizontal" else ImageOps.flip(img)
        buf = io.BytesIO()
        img.save(buf, format="PNG")
        img_doc = pymupdf.open("png", buf.getvalue())
        pdfbytes = img_doc.convert_to_pdf()
        img_doc.close()
        img_pdf = pymupdf.open("pdf", pdfbytes)
        npage = out.new_page(width=page.rect.width, height=page.rect.height)
        npage.show_pdf_page(npage.rect, img_pdf, 0)
        img_pdf.close()
    src.close()
    return _save(out, dest)


def n_up(path: str, dest: Path, options: dict) -> Path:
    n = int(options.get("n", 2))
    cols = 2 if n >= 2 else 1
    rows = math.ceil(n / cols)
    src = _open(path)
    out = pymupdf.open()
    w, h = 595, 842
    idx = 0
    while idx < src.page_count:
        page = out.new_page(width=w, height=h)
        cell_w, cell_h = w / cols, h / rows
        for r in range(rows):
            for c in range(cols):
                if idx >= src.page_count:
                    break
                rect = pymupdf.Rect(c * cell_w + 8, r * cell_h + 8, (c + 1) * cell_w - 8, (r + 1) * cell_h - 8)
                page.show_pdf_page(rect, src, idx)
                idx += 1
    src.close()
    return _save(out, dest)


def extract_images(path: str, dest_zip: Path) -> Path:
    doc = _open(path)
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        n = 0
        for i, page in enumerate(doc, 1):
            for img in page.get_images(full=True):
                xref = img[0]
                data = doc.extract_image(xref)
                ext = data.get("ext", "png")
                n += 1
                zf.writestr(f"page-{i}-img-{n}.{ext}", data["image"])
        if n == 0:
            zf.writestr("readme.txt", "No embedded images found.")
    dest_zip.parent.mkdir(parents=True, exist_ok=True)
    dest_zip.write_bytes(buf.getvalue())
    doc.close()
    return dest_zip


def pdf_to_images(path: str, dest_zip: Path, options: dict) -> Path:
    dpi = int(options.get("quality", 150))
    doc = _open(path)
    pages = parse_ranges(options.get("ranges", ""), doc.page_count)
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        for i in pages:
            pix = doc[i].get_pixmap(dpi=dpi, alpha=False)
            zf.writestr(f"page-{i + 1}.jpg", pix.tobytes("jpeg"))
    dest_zip.parent.mkdir(parents=True, exist_ok=True)
    dest_zip.write_bytes(buf.getvalue())
    doc.close()
    return dest_zip


def pdf_to_text(path: str, dest: Path, options: dict) -> Path:
    doc = _open(path)
    pages = parse_ranges(options.get("ranges", ""), doc.page_count)
    texts = []
    for i in pages:
        texts.append(doc[i].get_text("text"))
    dest.parent.mkdir(parents=True, exist_ok=True)
    dest.write_text("\n\n".join(texts), encoding="utf-8")
    doc.close()
    return dest


def pdf_to_word(path: str, dest: Path, options: dict) -> Path:
    from docx import Document

    doc = _open(path)
    word = Document()
    pages = parse_ranges(options.get("ranges", ""), doc.page_count)
    for i in pages:
        word.add_heading(f"Page {i + 1}", level=2)
        text = doc[i].get_text("text")
        for para in text.split("\n"):
            if para.strip():
                word.add_paragraph(para)
        word.add_page_break()
    dest.parent.mkdir(parents=True, exist_ok=True)
    word.save(dest)
    doc.close()
    return dest


def pdf_to_excel(path: str, dest: Path, options: dict) -> Path:
    from openpyxl import Workbook

    doc = _open(path)
    wb = Workbook()
    pages = parse_ranges(options.get("ranges", ""), doc.page_count)
    first = True
    for i in pages:
        ws = wb.active if first else wb.create_sheet(f"Page {i + 1}")
        if first:
            ws.title = f"Page {i + 1}"
            first = False
        tabs = doc[i].find_tables()
        wrote = False
        if tabs and tabs.tables:
            for t in tabs.tables:
                data = t.extract()
                for row in data:
                    ws.append([c if c is not None else "" for c in row])
                wrote = True
        if not wrote:
            for line in doc[i].get_text("text").split("\n"):
                if line.strip():
                    ws.append([line])
    dest.parent.mkdir(parents=True, exist_ok=True)
    wb.save(dest)
    doc.close()
    return dest


def pdf_to_ppt(path: str, dest: Path, options: dict) -> Path:
    from pptx import Presentation
    from pptx.util import Inches, Emu

    src = _open(path)
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]
    pages = parse_ranges(options.get("ranges", ""), src.page_count)
    tmp = dest.parent / "_slides"
    tmp.mkdir(exist_ok=True)
    for i in pages:
        slide = prs.slides.add_slide(blank)
        pix = src[i].get_pixmap(dpi=120, alpha=False)
        img_path = tmp / f"s{i}.jpg"
        pix.save(str(img_path))
        slide.shapes.add_picture(str(img_path), Emu(0), Emu(0), prs.slide_width, prs.slide_height)
    dest.parent.mkdir(parents=True, exist_ok=True)
    prs.save(dest)
    src.close()
    return dest


def images_to_pdf(paths: list[str], dest: Path, options: dict) -> Path:
    out = pymupdf.open()
    size_map = {"a4": (595, 842), "letter": (612, 792)}
    target = size_map.get(str(options.get("page_size", "a4")).lower(), (595, 842))
    for p in paths:
        if p.lower().endswith(".pdf"):
            src = pymupdf.open(p)
            out.insert_pdf(src)
            src.close()
            continue
        img = pymupdf.open(p)
        pdfbytes = img.convert_to_pdf()
        img.close()
        imgpdf = pymupdf.open("pdf", pdfbytes)
        page = out.new_page(width=target[0], height=target[1])
        page.show_pdf_page(page.rect, imgpdf, 0)
        imgpdf.close()
    return _save(out, dest)


def html_to_pdf(html: str, dest: Path, options: dict) -> Path:
    dest.parent.mkdir(parents=True, exist_ok=True)
    story = html or "<h1>Empty document</h1>"
    c = pdf_canvas.Canvas(str(dest), pagesize=letter)
    width, height = letter
    y = height - 72
    text = re.sub(r"<[^>]+>", " ", story)
    for line in text.split():
        c.drawString(72, y, line[:90])
        y -= 16
        if y < 72:
            c.showPage()
            y = height - 72
    c.save()
    try:
        doc = pymupdf.open()
        page = doc.new_page(width=595, height=842)
        page.insert_htmlbox(page.rect + (36, 36, -36, -36), story)
        return _save(doc, dest)
    except Exception:
        return dest


def word_to_pdf(path: str, dest: Path, options: dict) -> Path:
    from docx import Document

    document = Document(path)
    dest.parent.mkdir(parents=True, exist_ok=True)
    c = pdf_canvas.Canvas(str(dest), pagesize=letter)
    width, height = letter
    y = height - 72
    for para in document.paragraphs:
        text = para.text.strip()
        if not text:
            y -= 10
            continue
        for chunk in [text[i : i + 95] for i in range(0, len(text), 95)]:
            c.drawString(72, y, chunk)
            y -= 16
            if y < 72:
                c.showPage()
                y = height - 72
    c.save()
    return dest


def fill_sign(path: str, dest: Path, options: dict) -> Path:
    doc = _open(path)
    fields = options.get("fields") or []
    for field in fields:
        page_i = max(0, int(field.get("page", 1)) - 1)
        if page_i >= doc.page_count:
            continue
        page = doc[page_i]
        kind = field.get("type", "text")
        x = float(field.get("x", 72))
        y = float(field.get("y", 72))
        if kind == "text":
            page.insert_text((x, y), str(field.get("value", "")), fontsize=float(field.get("size", 12)), color=(0.05, 0.08, 0.16))
        elif kind == "check":
            page.insert_text((x, y), "X", fontsize=14, color=(0.1, 0.1, 0.15))
        elif kind == "date":
            page.insert_text((x, y), str(field.get("value", "")), fontsize=11, color=(0.1, 0.1, 0.15))
        elif kind == "signature":
            page.insert_text((x, y), str(field.get("value", "Signed")), fontsize=16, fontname="Times-Italic", color=(0.08, 0.12, 0.28))
    return _save(doc, dest)


def simple_edit(path: str, dest: Path, options: dict) -> Path:
    doc = _open(path)
    actions = options.get("actions") or []
    for act in actions:
        page_i = max(0, int(act.get("page", 1)) - 1)
        if page_i >= doc.page_count:
            continue
        page = doc[page_i]
        kind = act.get("type")
        x, y = float(act.get("x", 72)), float(act.get("y", 72))
        if kind == "text":
            page.insert_text((x, y), str(act.get("value", "")), fontsize=float(act.get("size", 14)), color=(0.08, 0.1, 0.18))
        elif kind == "highlight":
            w, h = float(act.get("w", 120)), float(act.get("h", 16))
            annot = page.add_highlight_annot(pymupdf.Rect(x, y, x + w, y + h))
            annot.set_colors(stroke=(1, 0.92, 0.3))
            annot.update()
        elif kind == "strikethrough":
            w = float(act.get("w", 120))
            annot = page.add_strikeout_annot(pymupdf.Rect(x, y, x + w, y + 12))
            annot.update()
        elif kind == "whiteout":
            w, h = float(act.get("w", 120)), float(act.get("h", 20))
            page.draw_rect(pymupdf.Rect(x, y, x + w, y + h), color=(1, 1, 1), fill=(1, 1, 1))
        elif kind == "rect":
            w, h = float(act.get("w", 80)), float(act.get("h", 40))
            page.draw_rect(pymupdf.Rect(x, y, x + w, y + h), color=(0.2, 0.3, 0.5), width=1.2)
        elif kind == "link":
            w, h = float(act.get("w", 120)), float(act.get("h", 18))
            page.insert_text((x, y + 12), str(act.get("value", "Link")), fontsize=11, color=(0.12, 0.32, 0.72))
            page.insert_link({"kind": pymupdf.LINK_URI, "from": pymupdf.Rect(x, y, x + w, y + h), "uri": str(act.get("url", "https://"))})
        elif kind == "replace":
            find = act.get("find") or ""
            repl = act.get("replace") or ""
            if find:
                hits = page.search_for(find)
                for rect in hits:
                    page.add_redact_annot(rect, text=repl, fontsize=11)
                page.apply_redactions()
    return _save(doc, dest)


def zip_paths(paths: list[Path], dest: Path) -> Path:
    dest.parent.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(dest, "w", zipfile.ZIP_DEFLATED) as zf:
        for p in paths:
            zf.write(p, arcname=p.name)
    return dest


def run_tool(tool: str, input_paths: list[str], dest: Path, options: dict[str, Any]) -> tuple[Path, list[Path]]:
    extra: list[Path] = []
    dest.parent.mkdir(parents=True, exist_ok=True)
    primary = input_paths[0] if input_paths else ""

    mapping = {
        "merge": lambda: merge(input_paths, dest, options),
        "split": lambda: None,
        "extract": lambda: extract_pages(primary, dest, options),
        "delete-pages": lambda: delete_pages(primary, dest, options),
        "rotate": lambda: rotate_pages(primary, dest, options),
        "crop": lambda: crop_pages(primary, dest, options),
        "resize": lambda: resize_pages(primary, dest, options),
        "organize": lambda: organize(primary, dest, options),
        "compress": lambda: compress(primary, dest, options),
        "ocr": lambda: ocr_pdf(primary, dest, options),
        "deskew": lambda: deskew(primary, dest, options),
        "protect": lambda: protect(primary, dest, options),
        "unlock": lambda: unlock(primary, dest, options),
        "watermark": lambda: watermark(primary, dest, options),
        "flatten": lambda: flatten_pdf(primary, dest, options),
        "metadata": lambda: edit_metadata(primary, dest, options),
        "bookmarks": lambda: create_bookmarks(primary, dest, options),
        "header-footer": lambda: header_footer(primary, dest, options),
        "page-numbers": lambda: header_footer(primary, dest, {**options, "page_numbers": True}),
        "remove-annotations": lambda: remove_annotations(primary, dest, options),
        "grayscale": lambda: grayscale(primary, dest, options),
        "flip": lambda: flip_pages(primary, dest, options),
        "n-up": lambda: n_up(primary, dest, options),
        "extract-images": lambda: extract_images(primary, dest.with_suffix(".zip")),
        "pdf-to-jpg": lambda: pdf_to_images(primary, dest.with_suffix(".zip"), options),
        "pdf-to-text": lambda: pdf_to_text(primary, dest.with_suffix(".txt"), options),
        "pdf-to-word": lambda: pdf_to_word(primary, dest.with_suffix(".docx"), options),
        "pdf-to-excel": lambda: pdf_to_excel(primary, dest.with_suffix(".xlsx"), options),
        "pdf-to-ppt": lambda: pdf_to_ppt(primary, dest.with_suffix(".pptx"), options),
        "jpg-to-pdf": lambda: images_to_pdf(input_paths, dest, options),
        "html-to-pdf": lambda: html_to_pdf(options.get("html", ""), dest, options),
        "word-to-pdf": lambda: word_to_pdf(primary, dest, options),
        "fill-sign": lambda: fill_sign(primary, dest, options),
        "editor": lambda: simple_edit(primary, dest, options),
        "repair": lambda: _save(_open(primary), dest),
        "rename": lambda: _save(_open(primary), dest),
        "alternate-mix": lambda: merge(input_paths, dest, options),
        "split-half": lambda: None,
        "split-size": lambda: None,
        "split-bookmarks": lambda: None,
        "split-text": lambda: None,
        "batch-numbering": lambda: header_footer(primary, dest, {**options, "page_numbers": True}),
        "create-forms": lambda: simple_edit(primary, dest, options),
    }

    if tool in ("split", "split-pages"):
        parts = split_by_pages(primary, dest.parent / f"{dest.stem}-parts", options)
        extra = parts
        return zip_paths(parts, dest.with_suffix(".zip")), extra
    if tool == "split-half":
        parts = split_in_half(primary, dest.parent / f"{dest.stem}-parts")
        return zip_paths(parts, dest.with_suffix(".zip")), parts
    if tool == "split-size":
        parts = split_by_size(primary, dest.parent / f"{dest.stem}-parts", float(options.get("max_mb", 2)))
        return zip_paths(parts, dest.with_suffix(".zip")), parts
    if tool == "split-bookmarks":
        doc = _open(primary)
        toc = doc.get_toc()
        dest_dir = dest.parent / f"{dest.stem}-parts"
        dest_dir.mkdir(parents=True, exist_ok=True)
        pages = [t[2] - 1 for t in toc] or [0]
        pages.append(doc.page_count)
        parts = []
        for i in range(len(pages) - 1):
            start, end = pages[i], pages[i + 1] - 1
            if end < start:
                continue
            out = pymupdf.open()
            out.insert_pdf(doc, from_page=start, to_page=end)
            p = dest_dir / f"bookmark-{i + 1}.pdf"
            out.save(p)
            out.close()
            parts.append(p)
        doc.close()
        return zip_paths(parts or [_save(_open(primary), dest)], dest.with_suffix(".zip")), parts
    if tool == "split-text":
        needle = options.get("text") or ""
        doc = _open(primary)
        dest_dir = dest.parent / f"{dest.stem}-parts"
        dest_dir.mkdir(parents=True, exist_ok=True)
        cuts = [0]
        if needle:
            for i, page in enumerate(doc):
                if needle.lower() in page.get_text("text").lower():
                    cuts.append(i)
        cuts.append(doc.page_count)
        cuts = sorted(set(cuts))
        parts = []
        for i in range(len(cuts) - 1):
            start, end = cuts[i], cuts[i + 1] - 1
            if end < start:
                continue
            out = pymupdf.open()
            out.insert_pdf(doc, from_page=start, to_page=end)
            p = dest_dir / f"text-{i + 1}.pdf"
            out.save(p)
            out.close()
            parts.append(p)
        doc.close()
        return zip_paths(parts or [_save(_open(primary), dest)], dest.with_suffix(".zip")), parts

    fn = mapping.get(tool)
    if not fn:
        raise ValueError(f"Unknown tool: {tool}")
    result = fn()
    if result is None:
        raise ValueError(f"Tool {tool} produced no output")
    return result, extra
